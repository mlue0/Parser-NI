"""Экспорт данных формы в Excel (.xlsx) и PDF."""

from __future__ import annotations

from pathlib import Path

from logs.setup import get_logger

logger = get_logger(__name__)


# ── Excel ──────────────────────────────────────────────────────

def export_to_excel(
    labeled_values: list[tuple[str, str]],
    file_path: str | Path,
) -> None:
    """Сохраняет данные в .xlsx — две колонки: «Поле» / «Значение»."""
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError as exc:
        raise RuntimeError("Для экспорта в Excel нужен пакет openpyxl") from exc

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Данные разбраковки"

    # Заголовки
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1a3a5c")
    headers = ["Поле", "Значение"]
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    # Данные
    alt_fill = PatternFill("solid", fgColor="f0f4f8")
    for i, (label, value) in enumerate(labeled_values, start=2):
        ws.cell(row=i, column=1, value=label)
        ws.cell(row=i, column=2, value=value)
        if i % 2 == 0:
            for col in (1, 2):
                ws.cell(row=i, column=col).fill = alt_fill

    # Ширина колонок
    ws.column_dimensions["A"].width = 35
    ws.column_dimensions["B"].width = 30

    wb.save(file_path)
    logger.info("Экспорт в Excel: %s", file_path)


# ── PDF ────────────────────────────────────────────────────────

def export_to_pdf(
    labeled_values: list[tuple[str, str]],
    file_path: str | Path,
    title: str = "Данные разбраковки кристаллов",
) -> None:
    """Сохраняет данные в PDF через reportlab или html+QPrinter fallback."""
    try:
        _export_pdf_reportlab(labeled_values, file_path, title)
    except ImportError:
        _export_pdf_qprinter(labeled_values, file_path, title)


def _export_pdf_reportlab(
    labeled_values: list[tuple[str, str]],
    file_path: str | Path,
    title: str,
) -> None:
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.units import cm
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    except ImportError as exc:
        raise ImportError("reportlab не установлен") from exc

    doc = SimpleDocTemplate(str(file_path), pagesize=A4,
                            leftMargin=2 * cm, rightMargin=2 * cm,
                            topMargin=2 * cm, bottomMargin=2 * cm)
    styles = getSampleStyleSheet()
    story = []
    story.append(Paragraph(title, styles["Title"]))
    story.append(Spacer(1, 0.5 * cm))

    table_data = [["Поле", "Значение"]] + [[lbl, val] for lbl, val in labeled_values]
    tbl = Table(table_data, colWidths=[8 * cm, 9 * cm])
    tbl.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1a3a5c")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f0f4f8")]),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cccccc")),
        ("PADDING", (0, 0), (-1, -1), 6),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    story.append(tbl)
    doc.build(story)
    logger.info("Экспорт в PDF (reportlab): %s", file_path)


def _export_pdf_qprinter(
    labeled_values: list[tuple[str, str]],
    file_path: str | Path,
    title: str,
) -> None:
    """Fallback: генерация PDF через QTextDocument + QPrinter."""
    from html import escape

    from PyQt6.QtGui import QTextDocument
    from PyQt6.QtPrintSupport import QPrinter

    rows = "".join(
        f"<tr style='background:{'#f0f4f8' if i%2==0 else 'white'}'>"
        f"<td style='padding:6px;border:1px solid #ccc'>{escape(str(lbl))}</td>"
        f"<td style='padding:6px;border:1px solid #ccc'>{escape(str(val))}</td></tr>"
        for i, (lbl, val) in enumerate(labeled_values)
    )
    html = (
        f"<h2 style='color:#1a3a5c'>{escape(str(title))}</h2>"
        "<table style='border-collapse:collapse;width:100%'>"
        "<tr style='background:#1a3a5c;color:white'>"
        "<th style='padding:8px'>Поле</th><th style='padding:8px'>Значение</th></tr>"
        f"{rows}</table>"
    )
    doc = QTextDocument()
    doc.setHtml(html)

    printer = QPrinter(QPrinter.PrinterMode.HighResolution)
    printer.setOutputFormat(QPrinter.OutputFormat.PdfFormat)
    printer.setOutputFileName(str(file_path))
    printer.setPageSize(printer.pageLayout().pageSize())
    doc.print(printer)
    logger.info("Экспорт в PDF (QPrinter): %s", file_path)
